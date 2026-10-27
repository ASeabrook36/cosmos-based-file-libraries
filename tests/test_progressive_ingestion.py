"""Text-first publication, resumable enrichment and same-generation snapshot safety."""

import io

import pytest

from tests.fakes import FakeChat
from tests.test_images import office, picture
from tests.test_worker_handlers import ctx as worker_ctx
from tests.test_worker_handlers import doc, drain, replace, upload

ctx = worker_ctx


async def prepare(ctx):
    client, services, lib, worker = ctx
    worker.vision = FakeChat({"description": "GDP growth chart"})
    document = upload(client, lib, "outlook.docx", office("docx"))
    message = services.queue.sent.pop(0)[1]
    await worker.handle(message)
    return document, message


def read(client, lib, document):
    return client.post(
        f"/v1/libraries/{lib}/documents:read",
        headers={"X-Agent-Id": "agent"},
        json={"documents": [{"documentId": document}]},
    )


def search(client, lib, kind="text"):
    return client.post(
        f"/v1/libraries/{lib}/search",
        headers={"X-Agent-Id": "agent"},
        json={"query": "GDP growth productivity", "filters": {"contentKinds": [kind]}},
    ).json()


async def test_text_is_searchable_and_complete_before_vision_then_images_append_without_reembedding_text(ctx):
    client, services, lib, worker = ctx
    document, _ = await prepare(ctx)
    text_record = doc(client, lib, document)
    assert text_record["activeGeneration"] == text_record["pendingGeneration"] == 1
    assert text_record["imageStatus"] == "queued" and text_record["imagePendingCount"] == 1
    assert text_record["publishedChunkCount"] == text_record["chunkCount"]
    assert not worker.vision.calls
    assert search(client, lib)["results"] and not search(client, lib, "image")["results"]
    response = read(client, lib, document).json()
    assert response["complete"] and not response["documents"][0]["visualsComplete"]
    assert "productivity" in str(response) and any("image enrichment" in note for note in response["notes"])
    before = {k: dict(v) for k, v in services.chunks.items.items()}
    await drain(services, worker)
    enriched = doc(client, lib, document)
    assert enriched["contentRevision"] == 2 and enriched["imageStatus"] == "ready"
    assert enriched["activeGeneration"] == 1 and enriched["pendingGeneration"] is None
    assert enriched["chunkCount"] == text_record["chunkCount"] + 1
    assert len(worker.vision.calls) == 1 and read(client, lib, document).json()["documents"][0]["visualsComplete"]
    assert enriched["fullText"]["blobPath"] != text_record["fullText"]["blobPath"]
    assert all(services.chunks.items[k] == value for k, value in before.items())
    assert search(client, lib, "image")["results"]


async def test_other_uploads_in_same_library_publish_text_before_queued_images(ctx):
    client, services, lib, worker = ctx
    worker.vision = FakeChat({"description": "chart"})
    first = upload(client, lib, "first.docx", office("docx"))
    second = upload(client, lib, "second.docx", office("docx"))
    await worker.handle(services.queue.sent.pop(0)[1])
    await worker.handle(services.queue.sent.pop(0)[1])
    assert doc(client, lib, first)["activeGeneration"] == doc(client, lib, second)["activeGeneration"] == 1
    assert not worker.vision.calls and all(m["type"] == "enrichImages" for _, m in services.queue.sent)
    await drain(services, worker)
    assert doc(client, lib, first)["imageStatus"] == doc(client, lib, second)["imageStatus"] == "ready"


async def test_enrichment_loads_one_stored_crop_at_a_time_without_uploading_it_again(ctx, monkeypatch):
    from docx import Document

    client, services, lib, worker = ctx
    source = Document()
    source.add_heading("Regional charts", 1)
    source.add_paragraph("Compare regional growth and productivity using these charts.")
    for _ in range(3):
        source.add_picture(io.BytesIO(picture()))
    data = io.BytesIO()
    source.save(data)
    worker.vision = FakeChat({"description": "Regional GDP chart"})
    document = upload(client, lib, "charts.docx", data.getvalue())
    await worker.handle(services.queue.sent.pop(0)[1])
    prefix = f"{lib}/{document}/images/1/"
    downloaded = []
    original_download, original_create = services.blobs.download, services.blobs.create

    async def download(path):
        if path.startswith(prefix):
            # The previous crop must already have been described before another is loaded.
            assert len(worker.vision.calls) == len(downloaded)
            downloaded.append(path)
        return await original_download(path)

    async def create(path, data, content_type):
        assert not path.startswith(prefix), "prepared crops must not be uploaded again during enrichment"
        return await original_create(path, data, content_type)

    monkeypatch.setattr(services.blobs, "download", download)
    monkeypatch.setattr(services.blobs, "create", create)
    await drain(services, worker)
    assert len(downloaded) == len(worker.vision.calls) == doc(client, lib, document)["imageCount"] == 3


async def test_final_image_failure_keeps_text_and_reports_incomplete_visual_coverage(ctx):
    client, services, lib, worker = ctx
    document, _ = await prepare(ctx)
    message = services.queue.sent.pop(0)[1]
    worker.vision = FakeChat(RuntimeError("vision unavailable"))
    with pytest.raises(RuntimeError):
        await worker.handle(message)
    await worker.fail(message, "vision unavailable")
    await drain(services, worker)
    current = doc(client, lib, document)
    assert current["status"] == "ready" and current["imageStatus"] == "failed"
    assert current["activeGeneration"] == 1 and current["pendingGeneration"] is None
    assert search(client, lib)["results"]
    response = read(client, lib, document).json()
    assert response["complete"] and not response["documents"][0]["visualsComplete"]
    assert any("enrichment failed" in n for n in search(client, lib)["coverage"]["notes"])


async def test_failed_enrichment_enqueue_retries_without_reembedding_text(ctx, monkeypatch):
    client, services, lib, worker = ctx
    real_send = services.queue.send

    async def fail_image_enqueue(message, *, session_id):
        if message["type"] == "enrichImages":
            raise RuntimeError("queue unavailable")
        await real_send(message, session_id=session_id)

    worker.vision = FakeChat({"description": "chart"})
    document = upload(client, lib, "chart.docx", office("docx"))
    message = services.queue.sent.pop(0)[1]
    monkeypatch.setattr(services.queue, "send", fail_image_enqueue)
    with pytest.raises(RuntimeError):
        await worker.handle(message)
    assert doc(client, lib, document)["activeGeneration"] == 1
    calls = worker.embedder.calls
    monkeypatch.setattr(services.queue, "send", real_send)
    await worker.handle(message)
    assert worker.embedder.calls == calls
    await drain(services, worker)
    assert doc(client, lib, document)["imageStatus"] == "ready"


async def test_written_image_chunks_remain_hidden_until_enriched_snapshot_commits(ctx, monkeypatch):
    client, services, lib, worker = ctx
    document, _ = await prepare(ctx)
    real_upsert = services.chunks.upsert
    checked = []

    async def check_visibility(chunk):
        result = await real_upsert(chunk)
        if chunk["contentKind"] == "image":
            checked.append(True)
            assert not search(client, lib, "image")["results"]
            assert not read(client, lib, document).json()["documents"][0]["visualsComplete"]
        return result

    monkeypatch.setattr(services.chunks, "upsert", check_visibility)
    await drain(services, worker)
    assert checked and search(client, lib, "image")["results"]


async def test_enrichment_during_complete_read_refuses_mixed_snapshot(ctx, monkeypatch):
    client, services, lib, worker = ctx
    document, _ = await prepare(ctx)
    message = services.queue.sent.pop(0)[1]
    real_download = services.blobs.download

    async def changing_download(path):
        result = await real_download(path)
        if path.endswith("/text-sections.json"):
            await worker.handle(message)
        return result

    monkeypatch.setattr(services.blobs, "download", changing_download)
    response = read(client, lib, document)
    assert response.status_code == 409 and response.json()["error"]["code"] == "document_changed"


async def test_replace_and_delete_between_phases_cannot_publish_obsolete_images(ctx):
    client, services, lib, worker = ctx
    document, _ = await prepare(ctx)
    replace(client, lib, document, "new.txt", b"Completely new text about employment.")
    await drain(services, worker)
    current = doc(client, lib, document)
    assert current["activeGeneration"] == 2 and current["imageCount"] == 0 and not worker.vision.calls
    assert not search(client, lib, "image")["results"]
    second, _ = await prepare(ctx)
    client.delete(f"/v1/libraries/{lib}/documents/{second}")
    await drain(services, worker)
    assert await services.documents.get(lib, second) is None
    assert not any(f"/{second}/" in p for p in services.blobs.blobs)


async def test_metadata_edit_between_phases_and_duplicate_ingest_preserve_final_snapshot(ctx):
    client, services, lib, worker = ctx
    document, message = await prepare(ctx)
    client.patch(f"/v1/libraries/{lib}/documents/{document}", json={"author": "Reviewed Author"})
    await drain(services, worker)
    final = doc(client, lib, document)
    await worker.handle(message)
    assert doc(client, lib, document)["fullText"] == final["fullText"]
    assert all(c["author"] == "Reviewed Author" for c in services.chunks.items.values())
    assert len(worker.vision.calls) == 1
