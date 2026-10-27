"""Issue #2: evidence projection, context accounting, source order and document grouping."""

import asyncio
import copy
import io
import json

import pytest

from fli.chunking import default_tokenizer
from fli.tool import tool_definitions
from fli_demo.agent import Reply, ToolCall, compact, model_view, run_turn
from tests.fakes import FakeChat
from tests.layout_fakes import FakeLayout
from tests.test_demo import scripted
from tests.test_images import office, vector_pdf
from tests.test_progressive_ingestion import read
from tests.test_worker_handlers import ctx as worker_ctx
from tests.test_worker_handlers import doc, drain, upload

ctx = worker_ctx


class EvidenceIndex:
    def __init__(self):
        self.full = {
            "library": {"name": "Evidence"},
            "complete": True,
            "disposition": "complete",
            "estimatedTokens": 20000,
            "tokenizer": "test",
            "usage": {"requestCharge": 5},
            "documents": [
                {
                    "documentId": "d",
                    "generation": 1,
                    "fileName": "evidence.pdf",
                    "complete": True,
                    "sections": [
                        {"text": "Evidence. " * 4000, "heading": "Findings", "location": {"page": 1}},
                        {
                            "text": "Chart description",
                            "contentKind": "image",
                            "imageId": "image-1",
                            "location": {"page": 1, "figure": "1.1"},
                            "image": {
                                "url": "https://example.invalid/chart.png?sig=private",
                                "expiresAt": "tomorrow",
                                "regions": [{"polygon": list(range(8))}] * 50,
                                "captionLocation": [{"polygon": list(range(8))}],
                                "width": 640,
                                "height": 400,
                                "mime": "image/png",
                                "kind": "pdf_figure",
                                "caption": "Figure 1: Growth.",
                                "captionStatus": "identified",
                                "captionAssociation": "adjacent_inferred",
                            },
                        },
                    ],
                }
            ],
            "notes": ["Extracted sections"],
        }
        self.row = {
            **self.full["documents"][0]["sections"][1],
            "fileName": "evidence.pdf",
            "documentId": "d",
            "generation": 1,
            "documentTokens": 20000,
            "representation": "extracted_sections",
        }
        self.options = []

    def tools(self, *args):
        return tool_definitions({"name": "Evidence"})

    def search(self, *args, **kwargs):
        self.options.append(kwargs)
        return {"library": {"name": "Evidence"}, "coverage": {"disposition": "thin"}, "results": [self.row]}

    def read_documents(self, *args, **kwargs):
        return copy.deepcopy(self.full)

    def read_images(self, *args, **kwargs):
        return {
            "library": {"name": "Evidence"},
            "results": [self.row],
            "pageSections": [{"text": "Page text"}],
            "complete": True,
            "notes": ["Page candidates"],
        }


@pytest.mark.parametrize(
    "operation,arguments",
    [
        ("search_library_", {"query": "growth", "libraryId": "injected"}),
        ("read_documents_library_", {"documents": [{"documentId": "d"}]}),
        ("inspect_images_library_", {"documentId": "d", "pages": [1]}),
    ],
)
def test_projected_tools_preserve_evidence_and_native_attachments(operation, arguments):
    index = EvidenceIndex()
    name = next(t["function"]["name"] for t in index.tools() if t["function"]["name"].startswith(operation))
    chat = scripted(Reply(None, [ToolCall("c", name, json.dumps(arguments))]), Reply("Found it"))
    turn = run_turn(chat, index, {"id": "a", "libraryIds": ["bound"]}, [], "Read evidence", context_token_budget=64000)
    content = next(m["content"] for m in turn.messages if m["role"] == "tool")
    for forbidden in ("sig=", '"url"', '"regions"', '"captionLocation"', '"usage"', '"tokenizer"'):
        assert forbidden not in content
    assert "Figure 1: Growth." in content and "adjacent_inferred" in content
    visuals = [
        p
        for m in turn.messages
        if isinstance(m.get("content"), list)
        for p in m["content"]
        if p.get("type") == "image_url"
    ]
    if operation == "read_documents_library_":
        assert not visuals  # candidate attachments do not imply model inspection
    else:
        assert len(visuals) == 1 and "sig=private" in visuals[0]["image_url"]["url"]
    assert turn.searches[0].images[0]["documentId"] == "d"
    if operation == "read_documents_library_":
        projected = json.loads(content)
        assert [s["text"] for s in projected["documents"][0]["sections"]] == [
            s["text"] for s in index.full["documents"][0]["sections"]
        ]
    if operation == "search_library_":
        assert index.options == [{"filters": None}]
        assert json.loads(content)["results"][0]["documentTokens"] == 20000


def test_projection_saves_tokens_on_image_bearing_fixture():
    index, tokenizer = EvidenceIndex(), default_tokenizer()
    raw = len(tokenizer.encode(json.dumps(index.full)))
    projected = len(tokenizer.encode(json.dumps(model_view(index.full))))
    assert projected < raw - 1000
    assert compact(index.search())["results"][0]["imageId"] == "image-1"


@pytest.mark.parametrize("reported_usage,expect_complete", [(1200, True), (None, False), (24000, False)])
def test_provider_prompt_usage_replaces_consumed_image_reservation(reported_usage, expect_complete):
    index = EvidenceIndex()
    calls = 0

    def chat(messages, tools):
        nonlocal calls
        calls += 1
        usage = {"prompt_tokens": reported_usage} if reported_usage is not None else {}
        if calls == 1:
            return Reply(None, [ToolCall("s", "search_library_evidence", '{"query":"chart"}')])
        if calls == 2:
            assert any(isinstance(m.get("content"), list) for m in messages)
            return Reply(
                None,
                [ToolCall("r", "read_documents_library_evidence", '{"documents":[{"documentId":"d"}]}')],
                usage,
            )
        return Reply("Done")

    turn = run_turn(
        chat, index, {"id": "a", "libraryIds": ["bound"]}, [], "Read it", context_token_budget=28000, image_tokens=12000
    )
    content = json.loads(next(m["content"] for m in turn.messages if m.get("tool_call_id") == "r"))
    assert content.get("complete") is expect_complete


@pytest.mark.parametrize("extension", ["pdf", "docx", "pptx"])
async def test_enriched_snapshot_interleaves_figures_without_changing_text_snapshot(ctx, extension):
    client, services, lib, worker = ctx
    worker.vision = FakeChat({"description": "Growth chart"})
    if extension == "pdf":
        worker.layout = FakeLayout()
        source = vector_pdf()
    else:
        source = office(extension)
        output = io.BytesIO()
        if extension == "docx":
            from docx import Document

            document = Document(io.BytesIO(source))
            document.add_heading("Later topic", 1)
            document.add_paragraph("Separate later explanation.")
        else:
            from pptx import Presentation

            document = Presentation(io.BytesIO(source))
            slide = document.slides.add_slide(document.slide_layouts[1])
            slide.shapes.title.text = "Later topic"
            slide.placeholders[1].text = "Separate later explanation."
        document.save(output)
        source = output.getvalue()
    document_id = upload(client, lib, f"source.{extension}", source)
    await worker.handle(services.queue.sent.pop(0)[1])
    initial = doc(client, lib, document_id)
    path = initial["fullText"]["blobPath"]
    text_bytes = await services.blobs.download(path)
    text_chunks = copy.deepcopy(services.chunks.items)
    await drain(services, worker)
    assert await services.blobs.download(path) == text_bytes
    assert all(services.chunks.items[k] == value for k, value in text_chunks.items())
    full = read(client, lib, document_id).json()
    assert full["complete"]
    sections = full["documents"][0]["sections"]
    image_index = next(i for i, section in enumerate(sections) if section.get("image"))
    assert image_index < len(sections) - 1
    assert [s["text"] for s in sections if not s.get("image")] == [s["text"] for s in json.loads(text_bytes)]
    if extension == "docx":
        assert [s["location"]["paragraph"] for s in sections] == [2, 3, 6]


async def test_concurrent_failure_cancels_other_descriptions_without_publishing(ctx, monkeypatch):
    from docx import Document

    from fli.worker import handlers
    from tests.test_images import picture

    client, services, lib, worker = ctx
    worker.vision = FakeChat({"description": "chart"})
    source = Document()
    source.add_paragraph("Text remains available.")
    for _ in range(2):
        source.add_picture(io.BytesIO(picture()))
    output = io.BytesIO()
    source.save(output)
    document_id = upload(client, lib, "failure.docx", output.getvalue())
    await worker.handle(services.queue.sent.pop(0)[1])
    before = copy.deepcopy(services.chunks.items)
    sibling_started, cancelled = asyncio.Event(), asyncio.Event()

    async def failing(model, visual, *args):
        if visual.location["paragraph"] == 2:
            await sibling_started.wait()
            raise RuntimeError("provider failed")
        sibling_started.set()
        try:
            await asyncio.Future()
        finally:
            cancelled.set()

    monkeypatch.setattr(handlers, "describe_cached", failing)
    with pytest.raises(RuntimeError, match="provider failed"):
        await asyncio.wait_for(worker.handle(services.queue.sent.pop(0)[1]), timeout=2)
    assert cancelled.is_set() and services.chunks.items == before
    assert read(client, lib, document_id).json()["complete"]


@pytest.mark.parametrize("file_type,key", [("pdf", "page"), ("pptx", "slide"), ("docx", "paragraph")])
def test_figure_order_handles_empty_regions_ties_and_legacy_locations(file_type, key):
    from fli.worker.handlers import interleave_sections

    text = [{"text": "first", "location": {key: 2}}, {"text": "second", "location": {key: 4}}]
    figures = [
        {"text": label, "location": {key: position}}
        for label, position in [("early", 1), ("a", 2), ("b", 2), ("gap", 3), ("last", 5)]
    ]
    assert [s["text"] for s in interleave_sections(text, figures, file_type)] == [
        "early",
        "first",
        "a",
        "b",
        "gap",
        "second",
        "last",
    ]
    legacy = [{"text": "unlocated", "location": {}}]
    assert interleave_sections(legacy, figures, file_type)[0] == legacy[0]


def test_reported_usage_allows_new_images_without_charging_old_images_twice():
    index = EvidenceIndex()
    index.full["documents"][0]["sections"][0]["text"] = "Brief explanation"
    index.full["documents"][0]["sections"][1]["imageId"] = "image-2"
    original_read = index.read_images

    def second_image(*args, **kwargs):
        response = original_read(*args, **kwargs)
        response["results"] = [{**response["results"][0], "imageId": "image-2"}]
        return response

    index.read_images = second_image
    chat = scripted(
        Reply(None, [ToolCall("s", "search_library_evidence", '{"query":"chart"}')]),
        Reply(
            None,
            [ToolCall("r", "inspect_images_library_evidence", '{"documentId":"d","imageIds":["image-2"]}')],
            {"prompt_tokens": 1200},
        ),
        Reply("Both figures received"),
    )
    turn = run_turn(
        chat, index, {"id": "a", "libraryIds": ["bound"]}, [], "Read it", context_token_budget=28000, image_tokens=12000
    )
    visuals = [
        p
        for m in turn.messages
        if isinstance(m.get("content"), list)
        for p in m["content"]
        if p.get("type") == "image_url"
    ]
    assert len(visuals) == 2
