"""Visual evidence from ingestion through search, full reads, and multimodal chat."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime

import pytest
from PIL import Image, ImageDraw

from fli.images import collect_visuals, describe
from tests.fakes import FakeChat
from tests.layout_fakes import FakeLayout
from tests.test_worker_handlers import ctx as worker_ctx  # noqa: F401
from tests.test_worker_handlers import doc, drain, replace, upload

ctx = worker_ctx


def picture():
    image = Image.new("RGB", (640, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 30), "GDP growth over time, 2020-2025", fill="black")
    draw.line([(40, 320), (600, 320)], fill="black", width=3)
    draw.line([(40, 320), (140, 280), (240, 290), (400, 190), (600, 80)], fill="blue", width=5)
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def office(kind):
    out = io.BytesIO()
    if kind == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading("Economic outlook", 1)
        doc.add_paragraph("Investment and productivity explain the observed expansion.")
        doc.add_picture(io.BytesIO(picture()))
        doc.add_paragraph("Figure 1: GDP growth between 2020 and 2025.")
    else:
        from pptx import Presentation
        from pptx.util import Inches

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[5])
        slide.shapes.title.text = "Economic outlook"
        slide.shapes.add_picture(io.BytesIO(picture()), Inches(1), Inches(2))
    doc.save(out)
    return out.getvalue()


def vector_pdf():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 16 Tf 50 700 Td (GDP growth 2020-2025) Tj ET\n"
        b"0 0 1 RG 3 w 50 100 m 150 180 l 250 160 l 400 300 l 550 450 l S"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_office_images_keep_heading_location_and_context(kind):
    [visual], skipped = collect_visuals(office(kind), f"outlook.{kind}", min_bytes=0)
    assert not skipped and (visual.width, visual.height) == (640, 400)
    assert visual.heading == "Economic outlook"
    assert visual.location == ({"paragraph": 3} if kind == "docx" else {"slide": 1})
    if kind == "docx":
        assert "productivity" in visual.context and "Figure 1" in visual.context
        assert visual.caption == "Figure 1: GDP growth between 2020 and 2025."
        assert visual.caption_location == [{"paragraph": 4}]


def test_pdf_primitives_cannot_be_promoted_to_page_images():
    with pytest.raises(ValueError, match="whole-page fallback is disabled"):
        collect_visuals(vector_pdf(), "outlook.pdf")


def test_skip_rules_cap_and_unsupported_media_are_counted():
    assert collect_visuals(picture(), "chart.png")[1] == 0  # compression does not discard a readable chart
    assert collect_visuals(picture(), "chart.png", min_bytes=10240)[1] == 1  # explicit operator override
    assert collect_visuals(picture(), "chart.png", min_bytes=0, min_pixels=1_000_000)[1] == 1
    assert collect_visuals(office("docx"), "outlook.docx", min_bytes=0, cap=0) == ([], 1)
    assert collect_visuals(b"bad data", "image.png", min_bytes=0) == ([], 1)


async def test_vision_prompt_includes_pixels_and_untrusted_source_context():
    [visual], _ = collect_visuals(office("docx"), "outlook.docx", min_bytes=0)
    model = FakeChat({"description": "A line chart of GDP growth over 2020 to 2025, trending upward."})
    text = await describe(model, visual, "outlook.docx")
    assert "GDP growth" in text and "productivity" in text and "paragraph 3" in text
    parts = model.calls[0][1]["content"]
    assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")
    assert "untrusted" in model.calls[0][0]["content"]


@pytest.mark.parametrize("extension", ["png", "docx", "pptx", "pdf"])
async def test_visual_ingestion_search_and_full_read_return_same_image(ctx, extension):
    client, services, lib, worker = ctx
    worker.vision = FakeChat({"description": "GDP growth line chart over time, 2020-2025, an upward trend."})
    worker.image_min_bytes = 0
    if extension == "pdf":
        worker.layout = FakeLayout()
    data = picture() if extension == "png" else vector_pdf() if extension == "pdf" else office(extension)
    document = upload(client, lib, f"outlook.{extension}", data=data)
    await drain(services, worker)
    stored = doc(client, lib, document)
    assert stored["status"] == "ready" and stored["imageCount"] == 1
    response = client.post(
        f"/v1/libraries/{lib}/search",
        headers={"X-Agent-Id": "agent"},
        json={"query": "I recall a chart about GDP over time", "filters": {"contentKinds": ["image"]}},
    )
    assert response.status_code == 200, response.text
    [hit] = response.json()["results"]
    assert hit["contentKind"] == "image" and "GDP" in hit["text"] and hit["imageId"] == hit["chunkId"]
    assert "imageRef" not in response.text
    assert datetime.fromisoformat(hit["image"]["expiresAt"]) > datetime.now(UTC)
    full = client.post(
        f"/v1/libraries/{lib}/documents:read",
        headers={"X-Agent-Id": "agent"},
        json={"documents": [{"documentId": document}]},
    ).json()
    assert full["complete"]
    [image] = [s for s in full["documents"][0]["sections"] if s.get("image")]
    assert image["imageId"] == hit["imageId"] and image["image"]["url"] == hit["image"]["url"]
    assert "imageRef" not in json.dumps(full)
    refused = client.post(
        f"/v1/libraries/{lib}/documents:read",
        headers={"X-Agent-Id": "agent"},
        json={"documents": [{"documentId": document}], "maxTokens": 1},
    ).json()
    assert not refused["complete"] and "sections" not in refused["documents"][0]


async def test_missing_vision_and_failed_replacement_preserve_active_image(ctx):
    client, services, lib, worker = ctx
    worker.image_min_bytes = 0
    document = upload(client, lib, "chart.png", data=picture())
    await drain(services, worker)
    assert doc(client, lib, document)["statusDetail"] == "vision_unavailable"
    worker.vision = FakeChat({"description": "GDP chart"})
    replace(client, lib, document, "chart.png", data=picture())
    await drain(services, worker)
    assert doc(client, lib, document)["activeGeneration"] == 2
    worker.vision = FakeChat(RuntimeError("vision service down"))
    generation = replace(client, lib, document, "new.png", data=picture())
    message = services.queue.sent.pop(0)[1]
    with pytest.raises(RuntimeError):
        await worker.handle(message)
    assert doc(client, lib, document)["activeGeneration"] == 2
    await worker.fail(message, "vision_failure")
    await drain(services, worker)
    assert doc(client, lib, document)["activeGeneration"] == 2
    assert not any(f"/images/{generation}/" in p for p in services.blobs.blobs)
    client.delete(f"/v1/libraries/{lib}/documents/{document}")
    await drain(services, worker)
    assert not any(f"/{document}/" in p for p in services.blobs.blobs)


def test_demo_delivers_image_to_model_and_preserves_inline_reference():
    from fli_demo.agent import Reply, ToolCall, run_turn
    from tests.test_demo import AGENT, FakeIndex, scripted

    index = FakeIndex()
    original = index.search

    def search(*args, **kwargs):
        response = original(*args, **kwargs)
        response["results"][0].update(
            imageId="image-1",
            contentKind="image",
            image={"url": "https://files.example.invalid/chart.png", "expiresAt": "later"},
        )
        return response

    index.search = search
    chat = scripted(
        Reply(None, [ToolCall("call1", "search_library_finance", '{"query":"GDP chart"}')]),
        Reply("Here it is: ![GDP chart](fli-image:image-1)"),
    )
    turn = run_turn(chat, index, AGENT, [], "Find the GDP chart")
    assert turn.searches[0].images[0]["imageId"] == "image-1"
    messages = chat.seen[1][0]
    assert messages[-1]["role"] == "user"
    assert messages[-1]["content"][1]["type"] == "image_url"
    assert messages[-2]["role"] == "tool" and '"imageId": "image-1"' in messages[-2]["content"]
    assert "fli-image:image-1" in turn.answer


async def test_deleted_during_image_link_issuance_is_not_returned(ctx, monkeypatch):
    client, services, lib, worker = ctx
    worker.image_min_bytes = 0
    worker.vision = FakeChat({"description": "GDP chart"})
    document = upload(client, lib, "chart.png", data=picture())
    await drain(services, worker)
    original = services.blobs.read_link

    async def deleting_link(path):
        result = await original(path)
        current = await services.documents.get(lib, document)
        current["status"] = "deleting"
        await services.documents.replace(current)
        return result

    monkeypatch.setattr(services.blobs, "read_link", deleting_link)
    response = client.post(f"/v1/libraries/{lib}/search", headers={"X-Agent-Id": "agent"}, json={"query": "GDP chart"})
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "document_changed"
    assert "https://" not in response.text


async def test_link_cannot_escape_the_active_document_prefix(ctx):
    client, services, lib, worker = ctx
    worker.image_min_bytes = 0
    worker.vision = FakeChat({"description": "GDP chart"})
    document = upload(client, lib, "chart.png", data=picture())
    await drain(services, worker)
    chunk = next(iter(services.chunks.items.values()))
    chunk["imageRef"] = f"{lib}/another-document/images/1/secret.png"
    response = client.post(f"/v1/libraries/{lib}/search", headers={"X-Agent-Id": "agent"}, json={"query": "GDP"})
    assert response.status_code == 409 and "secret" not in response.text
    assert doc(client, lib, document)["activeGeneration"] == 1


async def test_image_full_read_keeps_personal_binding_and_cross_library_isolation(ctx):
    client, services, lib, worker = ctx
    worker.image_min_bytes = 0
    worker.vision = FakeChat({"description": "GDP chart"})
    personal = client.post(
        "/v1/libraries", json={"kind": "personal", "name": "Private", "boundAgentId": "owner"}
    ).json()["id"]
    document = upload(client, personal, "chart.png", data=picture())
    await drain(services, worker)
    response = client.post(f"/v1/libraries/{personal}/search", headers={"X-Agent-Id": "other"}, json={"query": "GDP"})
    assert response.status_code == 403 and "url" not in response.text
    response = client.post(
        f"/v1/libraries/{personal}/documents:read",
        headers={"X-Agent-Id": "other"},
        json={"documents": [{"documentId": document}]},
    )
    assert response.status_code == 403 and "url" not in response.text
    response = client.post(
        f"/v1/libraries/{lib}/documents:read",
        headers={"X-Agent-Id": "owner"},
        json={"documents": [{"documentId": document}]},
    )
    assert response.status_code == 404 and "url" not in response.text


async def test_user_delegation_link_is_blob_scoped_https_read_only():
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from urllib.parse import parse_qs, urlsplit

    from azure.storage.blob import UserDelegationKey

    from fli.blobs import AzureBlobs

    key = UserDelegationKey()
    key.signed_oid = "00000000-0000-0000-0000-000000000001"
    key.signed_tid = "00000000-0000-0000-0000-000000000002"
    key.signed_start = "2026-01-01T00:00:00Z"
    key.signed_expiry = "2027-01-01T00:00:00Z"
    key.signed_service = "b"
    key.signed_version = "2024-11-04"
    key.value = "dGVzdC1zaWduaW5nLWtleQ=="
    service = SimpleNamespace(get_user_delegation_key=AsyncMock(return_value=key))
    container = SimpleNamespace(
        account_name="example",
        container_name="files",
        get_blob_client=lambda path: SimpleNamespace(url="https://example.blob.core.windows.net/files/" + path),
    )
    blobs = AzureBlobs(container, service)
    first = await blobs.read_link("library/document/images/1/image-0.png")
    second = await blobs.read_link("library/document/images/1/image-1.png")
    assert service.get_user_delegation_key.await_count == 1  # delegation key cached, individual SAS scoped per blob
    params = parse_qs(urlsplit(first["url"]).query)
    assert params["sp"] == ["r"] and params["spr"] == ["https"] and params["sr"] == ["b"]
    assert "skoid" in params and first["url"] != second["url"]
    assert 14 * 60 < (datetime.fromisoformat(first["expiresAt"]) - datetime.now(UTC)).total_seconds() <= 15 * 60


def test_demo_text_only_model_still_receives_description_and_image_link():
    from fli_demo.agent import Reply, ToolCall, run_turn
    from tests.test_demo import AGENT, FakeIndex, scripted

    index = FakeIndex()
    original = index.search

    def search(*args, **kwargs):
        response = original(*args, **kwargs)
        response["results"][0].update(
            imageId="image-1", contentKind="image", image={"url": "https://example.invalid/i.png"}
        )
        return response

    index.search = search
    chat = scripted(Reply(None, [ToolCall("c", "search_library_finance", '{"query":"chart"}')]), Reply("Found it"))
    turn = run_turn(chat, index, AGENT, [], "Find the chart", image_inputs=False)
    assert not any(isinstance(m.get("content"), list) for m in turn.messages)
    assert '"imageId": "image-1"' in chat.seen[1][0][-1]["content"]
    assert turn.searches[0].images


async def test_successful_image_replacement_changes_ids_and_purges_old_bytes(ctx):
    client, services, lib, worker = ctx
    worker.image_min_bytes = 0
    worker.vision = FakeChat({"description": "GDP chart"})
    document = upload(client, lib, "chart.png", data=picture())
    await drain(services, worker)
    generation = replace(client, lib, document, "updated.png", data=picture())
    await drain(services, worker)
    response = client.post(
        f"/v1/libraries/{lib}/search", headers={"X-Agent-Id": "agent"}, json={"query": "GDP chart"}
    ).json()
    assert response["results"] and all(r["generation"] == generation for r in response["results"])
    assert all(f":{generation}:" in r["imageId"] for r in response["results"])
    assert not any("/images/1/" in p for p in services.blobs.blobs)
    replace(client, lib, document, "plain.txt", data=b"The former image has been replaced with plain source text.")
    await drain(services, worker)
    assert doc(client, lib, document)["imageCount"] == 0
    assert not any("/images/" in p for p in services.blobs.blobs)


def test_transparent_chart_uses_a_visible_white_background():
    image = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    image.putpixel((128, 128), (0, 0, 0, 255))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    [visual], _ = collect_visuals(buffer.getvalue(), "transparent.png", min_bytes=0)
    with Image.open(io.BytesIO(visual.data)) as output:
        assert output.getpixel((0, 0)) == (255, 255, 255)
        assert output.getpixel((128, 128)) == (0, 0, 0)


async def test_vector_only_pdf_without_text_or_ocr_still_becomes_visual_evidence(ctx):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    client, services, lib, worker = ctx
    reader = PdfReader(io.BytesIO(vector_pdf()))
    writer = PdfWriter()
    page = writer.add_page(reader.pages[0])
    stream = DecodedStreamObject()
    stream.set_data(b"0 0 1 RG 3 w 50 100 m 150 180 l 250 160 l 400 300 l S")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    worker.vision = FakeChat({"description": "An unlabeled blue line chart."})
    worker.layout = FakeLayout(caption=False, body=False)
    document = upload(client, lib, "lines.pdf", data=output.getvalue())
    await drain(services, worker)
    stored = doc(client, lib, document)
    assert stored["status"] == "ready" and stored["imageCount"] == 1
    assert stored["extraction"]["method"] == "document_intelligence_layout"


def test_pptx_caption_association_uses_adjacent_text_not_all_slide_text():
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(picture()), Inches(1), Inches(1), width=Inches(4), height=Inches(2.5))
    caption = slide.shapes.add_textbox(Inches(1), Inches(3.6), Inches(4), Inches(0.3))
    caption.text = "Figure 1: Annual GDP growth."
    far = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.3))
    far.text = "Figure 2: Some other comparison."
    out = io.BytesIO()
    deck.save(out)
    [visual], _ = collect_visuals(out.getvalue(), "chart.pptx")
    assert visual.caption == caption.text
    assert visual.caption_location == [{"slide": 1, "shape": caption.shape_id}]
    assert "Some other comparison" in visual.context
