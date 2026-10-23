"""Extraction (unreadable-content classification, OCR fallback, locations) and deterministic chunking."""

from __future__ import annotations

import io

import pytest

from fli.chunking import ApproxTokenizer, Chunker, chunking_version_for, split_sentences
from fli.enums import ExtractionMethod
from fli.extraction import OLE_MAGIC, Section, classify_unreadable, extract


def docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Orion programme", level=1)
    d.add_paragraph("The orion shared signal is reviewed quarterly.")
    d.add_heading("Risks", level=2)
    d.add_paragraph("Budget risk is low.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_bytes() -> bytes:
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Ship the orion signal in Q2."
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["item", "cost"])
    ws.append(["orion", 42])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf_with_text(text: str = "orion signal in a pdf") -> bytes:
    stream = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def blank_pdf(password: str | None = None, owner_only: bool = False) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password is not None:
        writer.encrypt(user_password="" if owner_only else password, owner_password=password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


async def test_docx_sections_follow_headings():
    result = await extract(docx_bytes(), filename="plan.docx")
    assert result.outcome == "ok" and result.method == ExtractionMethod.NATIVE_TEXT
    assert [s.heading for s in result.sections] == ["Orion programme", "Risks"]
    assert "orion shared signal" in result.sections[0].text


async def test_pptx_and_xlsx_are_extracted():
    deck = await extract(pptx_bytes(), filename="deck.pptx")
    assert deck.outcome == "ok" and deck.sections[0].heading == "Roadmap" and "orion" in deck.text
    sheet = await extract(xlsx_bytes(), filename="budget.xlsx")
    assert sheet.outcome == "ok" and sheet.sections[0].heading == "Budget" and "item: orion; cost: 42" in sheet.text


async def test_pdf_with_text_layer_is_native():
    result = await extract(pdf_with_text(), filename="a.pdf")
    assert result.outcome == "ok" and "orion signal" in result.text and result.page_count == 1


async def test_pdf_without_text_needs_ocr_and_uses_the_bounded_fallback():
    no_ocr = await extract(blank_pdf(), filename="scan.pdf")
    assert no_ocr.outcome == "needs_ocr" and no_ocr.unsupported_reason == "no_text_layer_and_ocr_disabled"

    class FakeOcr:
        calls = 0
        max_pages = 10

        async def read(self, data, content_type):  # noqa: ANN001
            FakeOcr.calls += 1
            return "scanned orion text", 1

    result = await extract(blank_pdf(), filename="scan.pdf", ocr=FakeOcr())
    assert (
        result.outcome == "ok" and result.ocr_applied and result.method == ExtractionMethod.DOCUMENT_INTELLIGENCE_READ
    )
    native = await extract(pdf_with_text(), filename="a.pdf", ocr=FakeOcr())
    assert native.ocr_applied is False and FakeOcr.calls == 1  # OCR only when no text layer


async def test_password_protected_pdf_is_categorical_and_owner_only_pdf_is_readable():
    locked = await extract(blank_pdf("secret"), filename="locked.pdf")
    assert locked.outcome == "categorical" and locked.categorical_reason == "password_protected"
    owner_only = await extract(blank_pdf("secret", owner_only=True), filename="owner.pdf")
    assert owner_only.outcome in ("ok", "needs_ocr")


def test_read_channel_distinguishes_rms_from_password_from_unknown():
    rms = OLE_MAGIC + b"\x00" * 64 + "DRMEncryptedDataSpace".encode("utf-16-le")
    password = OLE_MAGIC + b"\x00" * 64 + "StrongEncryptionDataSpace".encode("utf-16-le")
    unknown = OLE_MAGIC + b"\x00" * 64
    assert classify_unreadable(rms) == ("encrypted", None)
    assert classify_unreadable(password) == ("categorical", "password_protected")
    assert classify_unreadable(unknown) == ("encrypted", None)  # unreadable after a successful read → exclude
    assert classify_unreadable(b"%PDF-1.7 /Encrypt /Filter /MicrosoftIRMServices") == ("encrypted", None)


async def test_office_extension_with_protected_container_is_classified_not_retried():
    rms_docx = OLE_MAGIC + b"\x00" * 64 + "DRMEncryptedDataSpace".encode("utf-16-le")
    result = await extract(rms_docx, filename="secret.docx")
    assert result.outcome == "encrypted" and result.method == ExtractionMethod.UNSUPPORTED
    pw = OLE_MAGIC + b"\x00" * 64 + "StrongEncryptionDataSpace".encode("utf-16-le")
    assert (await extract(pw, filename="locked.xlsx")).categorical_reason == "password_protected"


async def test_unsupported_types_are_reported_not_guessed():
    result = await extract(b"\x00\x01binary", filename="model.bin")
    assert result.outcome == "unsupported" and result.unsupported_reason == "unsupported_file_type"


async def test_markdown_and_html():
    md = await extract(b"# Title\nintro text\n## Part\nbody text", filename="x.md")
    assert [s.heading for s in md.sections] == ["Title", "Part"]
    html = await extract(b"<html><script>x()</script><h1>Head</h1><p>orion para</p></html>", filename="page.aspx")
    assert "orion para" in html.text and "x()" not in html.text


# --- chunking ---------------------------------------------------------------------------

TOK = ApproxTokenizer()


def long_section(sentences: int = 120, heading: str | None = "Section") -> Section:
    body = " ".join(
        f"Sentence number {i} talks about the orion signal and its quarterly review cadence." for i in range(sentences)
    )
    return Section(heading, body)


def test_chunks_are_deterministic_bounded_and_overlapping():
    chunker = Chunker(TOK)
    first = chunker.chunk([long_section()])
    second = Chunker(ApproxTokenizer()).chunk([long_section()])
    assert [c.text for c in first] == [c.text for c in second]
    assert [c.index for c in first] == list(range(len(first)))
    assert len(first) > 3
    assert all(c.token_count <= 500 for c in first)
    for a, b in zip(first, first[1:], strict=False):
        tail = split_sentences(a.text.split("\n", 1)[1])[-1]
        assert tail in b.text  # at least one trailing sentence carried forward


def test_overlap_is_at_most_fifty_tokens_of_whole_sentences():
    chunker = Chunker(TOK)
    chunks = chunker.chunk([long_section()])
    for a, b in zip(chunks, chunks[1:], strict=False):
        a_sentences = split_sentences(a.text.split("\n", 1)[1])
        b_sentences = split_sentences(b.text.split("\n", 1)[1])
        shared = [s for s in b_sentences if s in a_sentences]
        assert shared and sum(chunker.count(s) for s in shared) <= 50


def test_sentences_are_never_cut_unless_longer_than_the_budget():
    chunks = Chunker(TOK).chunk([long_section()])
    for c in chunks:
        body = c.text.split("\n", 1)[1]
        assert body.startswith("Sentence number") and body.endswith("cadence.")


def test_oversized_sentence_is_split_on_token_boundaries():
    monster = Section(None, "word " * 1500)
    chunks = Chunker(TOK).chunk([monster])
    assert len(chunks) >= 3 and all(c.token_count <= 500 for c in chunks)


def test_an_over_long_heading_is_capped_and_its_text_opens_the_body():
    """Review finding 2: a heading over the target once made one chunk per body token."""
    chunker = Chunker(TOK)
    heading = " ".join(f"h{i}" for i in range(520))
    plain = chunker.chunk([long_section(heading=None)])
    capped = chunker.chunk([long_section(heading=heading)])
    assert len(capped) <= len(plain) + 3
    assert all(c.token_count <= 500 for c in capped)
    assert {c.heading for c in capped} == {" ".join(f"h{i}" for i in range(64))}
    assert sum(" h519" in c.text for c in capped) == 1  # the full heading is indexed once, not lost


def test_a_heading_within_the_cap_is_carried_whole():
    heading = " ".join(f"h{i}" for i in range(64))
    chunks = Chunker(TOK).chunk([long_section(heading=heading)])
    assert all(c.heading == heading and c.text.startswith(heading + "\n") for c in chunks)


def test_headings_are_carried_and_sections_never_merge():
    chunks = Chunker(TOK).chunk([long_section(60, "Alpha"), long_section(60, "Bravo")])
    assert all(c.text.startswith(("Alpha\n", "Bravo\n")) for c in chunks)
    assert not any(c.text.startswith("Alpha") and "Bravo" in c.text for c in chunks)


def test_email_rules():
    chunker = Chunker(TOK)
    short = chunker.chunk([Section(None, "Short note about the orion signal.")], "email")
    assert len(short) == 1
    body = " ".join(f"Line {i} of a long email about orion." for i in range(200))
    assert chunker.count(body) > 700
    assert len(chunker.chunk([Section(None, body)], "email")) > 1


def test_unknown_document_kind_is_refused():
    with pytest.raises(ValueError):
        Chunker(TOK).chunk([long_section(4)], "chat")  # type: ignore[arg-type]


def test_chunks_carry_their_section_location():
    sections = [
        Section("Intro", "First page text about orion.", {"page": 1}),
        Section("Intro", "Second page text about orion.", {"page": 2}),
    ]
    chunks = Chunker(TOK).chunk(sections)
    assert [c.location for c in chunks] == [{"page": 1}, {"page": 2}]


async def test_pdf_pptx_and_xlsx_sections_carry_locations():
    pdf = await extract(pdf_with_text(), filename="a.pdf")
    assert pdf.sections[0].location == {"page": 1}
    deck = await extract(pptx_bytes(), filename="deck.pptx")
    assert deck.sections[0].location == {"slide": 1}
    sheet = await extract(xlsx_bytes(), filename="budget.xlsx")
    assert sheet.sections[0].location == {"sheet": "Budget", "rowStart": 2, "rowEnd": 2}


async def test_pptx_speaker_notes_are_extracted():
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.notes_slide.notes_text_frame.text = "Mention the orion delay."
    buf = io.BytesIO()
    p.save(buf)
    deck = await extract(buf.getvalue(), filename="deck.pptx")
    assert "Speaker notes: Mention the orion delay." in deck.text


def test_chunking_version_names_the_tokenizer():
    assert chunking_version_for(TOK) == "chunking-v1.1+approx-wordpiece"

    class Cl100k:
        name = "cl100k_base"

    assert chunking_version_for(Cl100k()) == "chunking-v1.1"


def test_approx_tokenizer_round_trips():
    text = "Orion, signal!  Quarterly\nreview."
    assert TOK.decode(TOK.encode(text)) == text


@pytest.mark.parametrize(
    "text,expected", [("One. Two! Three?", 3), ("First para.\n\nSecond para", 2), ("e.g. lower case stays", 1)]
)
def test_sentence_splitter(text, expected):
    assert len(split_sentences(text)) == expected


async def test_csv_rows_are_header_value_lines_with_row_locations():
    rows = "year,country,gdp\n" + "\n".join(f"{2000 + i},Wales,{i}" for i in range(60))
    result = await extract(rows.encode(), filename="data/gdp.csv")
    assert result.outcome == "ok" and [s.location["rowStart"] for s in result.sections] == [2, 27, 52]
    assert result.sections[0].location == {"sheet": "gdp", "rowStart": 2, "rowEnd": 26}
    assert result.sections[-1].location["rowEnd"] == 61  # header is row 1, 60 data rows
    assert "year: 2000; country: Wales; gdp: 0" in result.sections[0].text
    assert all(s.heading == "gdp" for s in result.sections)


async def test_csv_handles_semicolons_bom_and_blank_cells():
    result = await extract("\ufeffa;b\n1;\n;2\n".encode(), filename="x.csv")
    assert result.sections[0].text.splitlines() == ["Row 2. a: 1", "Row 3. b: 2"]


EML = (
    b"From: Jane Doe <jane@example.org>\nTo: team@example.org\nSubject: Q4 margin\n"
    b"Date: Tue, 11 Mar 2025 09:30:00 +0100\nContent-Type: text/plain; charset=utf-8\n\nMargins fell in EMEA.\n"
)


async def test_eml_has_subject_heading_headers_and_authoring_metadata():
    result = await extract(EML, filename="m.eml")
    assert result.outcome == "ok" and result.sections[0].heading == "Q4 margin"
    assert "From: Jane Doe <jane@example.org>" in result.text and "Margins fell in EMEA." in result.text
    assert result.author == "Jane Doe"
    assert result.authored_at is not None and result.authored_at.isoformat() == "2025-03-11T08:30:00+00:00"


async def test_eml_html_only_body_and_attachment_names():
    raw = (
        b"From: a@example.org\nSubject: Report\nMIME-Version: 1.0\nContent-Type: multipart/mixed; boundary=b\n\n"
        b"--b\nContent-Type: text/html\n\n<p>Hello <b>there</b></p>\n"
        b'--b\nContent-Type: application/pdf\nContent-Disposition: attachment; filename="q4.pdf"\n\nxx\n--b--\n'
    )
    result = await extract(raw, filename="m.eml")
    assert "Hello there" in result.text and "Attachments: q4.pdf" in result.text and result.author == "a@example.org"


async def test_eml_with_bad_date_still_extracts():
    result = await extract(b"Subject: s\nDate: not a date\n\nbody", filename="m.eml")
    assert result.outcome == "ok" and result.authored_at is None


async def test_html_headings_become_sections_and_scripts_are_dropped():
    html = (
        b"<html><head><title>t</title></head><body><script>x()</script>"
        b"<p>Intro text.</p><h2>Methods</h2><p>We did it.</p></body></html>"
    )
    result = await extract(html, filename="a.html")
    assert [(s.heading, s.text) for s in result.sections] == [(None, "Intro text."), ("Methods", "We did it.")]


async def test_office_and_pdf_metadata_are_returned():
    import docx

    doc = docx.Document()
    doc.core_properties.author = "Sam Lee"
    doc.add_paragraph("hello")
    buf = io.BytesIO()
    doc.save(buf)
    result = await extract(buf.getvalue(), filename="a.docx")
    assert result.author == "Sam Lee" and result.authored_at is not None and result.authored_at.tzinfo is not None
    assert (await extract(pdf_with_text(), filename="a.pdf")).author is None  # no metadata: None, not ""


async def test_docx_tables_stay_under_their_heading_in_document_order():
    import docx

    doc = docx.Document()
    doc.add_heading("Costs", level=1)
    doc.add_paragraph("Summary line.")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate([["item", "cost"], ["orion", "42"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    doc.add_heading("Risks", level=1)
    doc.add_paragraph("None known.")
    buf = io.BytesIO()
    doc.save(buf)
    result = await extract(buf.getvalue(), filename="a.docx")
    assert [(s.heading, s.text) for s in result.sections] == [
        ("Costs", "Summary line.\nitem | cost\norion | 42"),
        ("Risks", "None known."),
    ]


async def test_row_numbers_are_the_sheets_own_when_blank_rows_exist():
    result = await extract(b"a,b\n\n1,2\n\n\n3,4\n", filename="x.csv")
    assert result.sections[0].text.splitlines() == ["Row 3. a: 1; b: 2", "Row 6. a: 3; b: 4"]
    assert result.sections[0].location == {"sheet": "x", "rowStart": 3, "rowEnd": 6}


async def test_html_without_a_closing_head_keeps_its_body_and_drops_the_title():
    result = await extract(b"<html><head><title>T</title><body><h1>H</h1><p>Body text.</p>", filename="a.html")
    assert [(s.heading, s.text) for s in result.sections] == [("H", "Body text.")]


async def test_eml_with_an_unknown_charset_is_read_not_called_encrypted():
    raw = b"From: a@example.org\nSubject: s\nContent-Type: text/plain; charset=x-unknown-8bit\n\nplain words\n"
    result = await extract(raw, filename="m.eml")
    assert result.outcome == "ok" and "plain words" in result.text


async def test_a_text_format_parser_crash_is_a_parser_failure_not_encryption():
    result = await extract(
        b"a,b\n" + b'"' + b"x" * 200_000 + b'"\n', filename="big-field.csv"
    )  # over csv's field limit
    assert result.outcome == "failed" and result.failure_reason == "parser:Error"


async def test_docx_merged_cells_appear_once_but_equal_neighbours_are_kept():
    import docx

    doc = docx.Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(0, 2).text = "Other"
    for c in range(3):
        table.cell(1, c).text = "Yes"
    buf = io.BytesIO()
    doc.save(buf)
    result = await extract(buf.getvalue(), filename="a.docx")
    assert result.sections[0].text == "Merged | Other\nYes | Yes | Yes"


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("D:20250311093000+01'00'", "2025-03-11T08:30:00+00:00"),
        ("D:20250311093000-05'30", "2025-03-11T15:00:00+00:00"),
        ("D:20250311093000Z", "2025-03-11T09:30:00+00:00"),
        ("D:2025", "2025-01-01T00:00:00+00:00"),
        ("garbage", None),
        ("D:20251399", None),
    ],
)
def test_pdf_dates_honour_their_zone(raw, expected):
    from fli.extraction import _pdf_date

    parsed = _pdf_date(raw)
    assert (parsed.isoformat() if parsed else None) == expected


async def test_parsing_runs_off_the_event_loop_thread(monkeypatch):
    import threading

    from fli import extraction

    seen = {}
    real = extraction._docx

    def spy(data):
        seen["thread"] = threading.current_thread()
        return real(data)

    monkeypatch.setattr(extraction, "_docx", spy)
    import docx

    doc = docx.Document()
    doc.add_paragraph("hello")
    buf = io.BytesIO()
    doc.save(buf)
    await extract(buf.getvalue(), filename="a.docx")
    assert seen["thread"] is not threading.main_thread()  # lock renewal keeps running meanwhile
