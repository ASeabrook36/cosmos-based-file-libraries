"""Fault regressions: concurrent attempts cannot change published generation evidence."""

from __future__ import annotations

import asyncio
import hashlib
import io
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest

from fli.blobs import AzureBlobs, MemoryBlobs
from fli.chunking import ApproxTokenizer, Chunker
from fli.chunks import CosmosChunks, ImmutableChunkConflict, MemoryChunks
from fli.governor import Governor
from fli.images import collect_visuals, describe_cached
from fli.proposals import AzureOpenAIChat, ChatReply
from fli.worker.handlers import Worker
from tests.fakes import FakeChat, FakeEmbedder, api
from tests.layout_fakes import FakeLayout
from tests.test_images import picture
from tests.test_worker_handlers import drain, upload


def setup(vision):
    client, services = api()
    lib = client.post("/v1/libraries", json={"kind": "shared", "name": "Review"}).json()["id"]
    worker = Worker(
        services.libraries,
        services.documents,
        services.chunks,
        services.blobs,
        services.queue,
        FakeEmbedder(),
        Chunker(ApproxTokenizer()),
        vision=vision,
    )
    return client, services, lib, worker


async def test_late_duplicate_cannot_overwrite_published_artifact_chunks_or_retagged_metadata():
    entered, released = asyncio.Event(), asyncio.Event()

    class DelayedVision:
        async def complete_json(self, messages):
            entered.set()
            await released.wait()
            return ChatReply({"description": "Late alternative description of the GDP chart"})

    client, services, lib, older = setup(DelayedVision())
    document = upload(client, lib, "chart.png", data=picture())
    message = services.queue.sent.pop(0)[1]
    newer = replace(older, vision=FakeChat({"description": "Canonical chart description"}))
    task = asyncio.create_task(older.handle(message))
    await entered.wait()
    try:
        await newer.handle(message)
        published = await services.documents.get(lib, document)
        descriptor = published["fullText"]
        before = dict(next(iter(services.chunks.items.values())))
        client.patch(f"/v1/libraries/{lib}/documents/{document}", json={"author": "Reviewed Author"})
        await newer.retag(lib, document, 1)
    finally:
        released.set()
        await task
    current = await services.documents.get(lib, document)
    assert current["fullText"] == descriptor
    assert hashlib.sha256(await services.blobs.download(descriptor["blobPath"])).hexdigest() == descriptor["sha256"]
    chunk = next(iter(services.chunks.items.values()))
    assert chunk["text"] == before["text"] and chunk["embedding"] == before["embedding"]
    assert chunk["createdAt"] == before["createdAt"] and chunk["author"] == "Reviewed Author"
    response = client.post(
        f"/v1/libraries/{lib}/documents:read",
        headers={"X-Agent-Id": "agent"},
        json={"documents": [{"documentId": document}]},
    )
    assert response.status_code == 200 and response.json()["complete"]


async def test_restart_after_embedding_failure_reuses_frozen_plan_without_vision():
    class Down(FakeEmbedder):
        async def embed(self, texts):
            raise RuntimeError("embedding unavailable")

    vision = FakeChat({"description": "GDP chart"})
    client, services, lib, worker = setup(vision)
    document = upload(client, lib, "chart.png", data=picture())
    message = services.queue.sent.pop(0)[1]
    with pytest.raises(RuntimeError):
        await replace(worker, embedder=Down()).handle(message)
    assert len(vision.calls) == 1
    restarted_vision = FakeChat(RuntimeError("must not describe again"))
    restarted = replace(worker, vision=restarted_vision)
    await restarted.handle(message)
    assert not restarted_vision.calls
    stored = await services.documents.get(lib, document)
    assert stored["status"] == "ready" and stored["processingMetrics"]["planCacheHit"]
    assert stored["processingMetrics"]["visionPromptTokens"] == 0  # no duplicate model charge


async def test_mid_document_vision_failure_reuses_completed_image_checkpoint():
    from docx import Document

    doc = Document()
    doc.add_heading("First chart", 1)
    doc.add_picture(io.BytesIO(picture()))
    doc.add_heading("Second chart", 1)
    doc.add_picture(io.BytesIO(picture()))
    data = io.BytesIO()
    doc.save(data)
    vision = FakeChat({"description": "First image"}, RuntimeError("vision unavailable"))
    client, services, lib, worker = setup(vision)
    document = upload(client, lib, "charts.docx", data=data.getvalue())
    message = services.queue.sent.pop(0)[1]
    with pytest.raises(RuntimeError):
        await worker.handle(message)
    assert len(vision.calls) == 2
    vision.answers = [{"description": "Second image"}]
    await worker.handle(message)
    assert len(vision.calls) == 3
    stored = await services.documents.get(lib, document)
    assert stored["imageCount"] == 2 and stored["processingMetrics"]["visionCacheHits"] == 1


async def test_vision_cache_includes_context_and_model_configuration():
    blobs = MemoryBlobs()
    [visual], _ = collect_visuals(picture(), "chart.png")
    vision = FakeChat({"description": "GDP chart"})
    vision.cache_identity = "deployment-version-a"
    prefix = "library/document/text/1/"
    first = await describe_cached(vision, visual, "chart.png", blobs, prefix)
    again = await describe_cached(vision, visual, "chart.png", blobs, prefix)
    assert first == again and len(vision.calls) == 1
    visual.context = "Different interpretation context"
    await describe_cached(vision, visual, "chart.png", blobs, prefix)
    vision.cache_identity = "deployment-version-b"
    await describe_cached(vision, visual, "chart.png", blobs, prefix)
    assert len(vision.calls) == 3


async def test_purge_removes_generation_plans_and_vision_checkpoints():
    client, services, lib, worker = setup(FakeChat({"description": "GDP chart"}))
    document = upload(client, lib, "chart.png", data=picture())
    await drain(services, worker)
    assert any(p.endswith("/plan.json") for p in services.blobs.blobs)
    assert any("/vision-" in p for p in services.blobs.blobs)
    client.put(
        f"/v1/libraries/{lib}/documents/{document}/content",
        files={"file": ("new.txt", b"Plain replacement.", "text/plain")},
    )
    await drain(services, worker)
    assert not any("/text/1/" in p or "/images/1/" in p for p in services.blobs.blobs)
    client.delete(f"/v1/libraries/{lib}/documents/{document}")
    await drain(services, worker)
    assert not any(f"/{document}/" in p for p in services.blobs.blobs)


async def test_blob_create_is_conditional_and_does_not_hide_other_errors():
    from azure.core.exceptions import ResourceExistsError

    container = SimpleNamespace(
        upload_blob=AsyncMock(side_effect=[None, ResourceExistsError("exists"), RuntimeError("down")])
    )
    blobs = AzureBlobs(container)
    assert await blobs.create("plan.json", b"first", "application/json")
    assert not await blobs.create("plan.json", b"second", "application/json")
    with pytest.raises(RuntimeError):
        await blobs.create("plan.json", b"third", "application/json")
    assert all(call.kwargs["overwrite"] is False for call in container.upload_blob.call_args_list)


async def test_cosmos_duplicate_uses_create_and_checks_hash_without_replacing_item():
    from azure.cosmos.exceptions import CosmosResourceExistsError

    record = {"id": "doc:1:0", "libraryId": "lib", "documentId": "doc", "contentHash": "canonical"}
    container = SimpleNamespace(
        create_item=AsyncMock(side_effect=CosmosResourceExistsError(message="exists")),
        read_item=AsyncMock(return_value=record),
        upsert_item=AsyncMock(),
    )
    chunks = CosmosChunks(container, Governor("writes", 2))
    await chunks.upsert(record)
    container.upsert_item.assert_not_called()
    container.read_item.assert_awaited_once()
    with pytest.raises(ImmutableChunkConflict):
        await chunks.upsert({**record, "contentHash": "different"})


async def test_memory_duplicate_does_not_revert_mutable_metadata():
    chunks = MemoryChunks()
    record = {"id": "doc:1:0", "libraryId": "lib", "documentId": "doc", "contentHash": "canonical", "author": "old"}
    await chunks.upsert(record)
    await chunks.patch("lib", "doc", "doc:1:0", {"author": "new"})
    await chunks.upsert(record)
    assert chunks.items[("lib", "doc", "doc:1:0")]["author"] == "new"
    with pytest.raises(ImmutableChunkConflict):
        await chunks.upsert({**record, "contentHash": "different"})


async def test_vision_transport_governor_caps_calls_across_documents():
    active, peak = 0, 0

    async def handler(request):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.001)
        active -= 1
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"description":"chart"}'}}]})

    async def token():
        return "test-token"

    gov = Governor("vision", 2)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
        model = AzureOpenAIChat("https://example.invalid", "vision", token, http, governor=gov)
        await asyncio.gather(*(model.complete_json([]) for _ in range(8)))
    assert peak == 2 and gov.requests == 8


def decorated_pdf():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    from tests.test_images import vector_pdf

    writer = PdfWriter()
    for _ in range(2):
        page = writer.add_blank_page(width=612, height=792)
        stream = DecodedStreamObject()
        stream.set_data(b"0.8 0.8 0.8 RG 1 w 30 30 550 730 re S")
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_page(PdfReader(io.BytesIO(vector_pdf())).pages[0])
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


async def test_page_borders_do_not_consume_the_visual_cap_before_a_layout_figure():
    from fli.figures import extract_pdf_figures

    layout = FakeLayout(figure_page=3)
    extraction = await extract_pdf_figures(decorated_pdf(), 3, layout, MemoryBlobs(), "text/", "images/", cap=1)
    [visual] = extraction.visuals
    assert visual.location == {"page": 3, "figure": "3.1"} and extraction.skipped == 0
    assert layout.figure_calls == ["3.1"]


async def test_pdf_layout_handles_vector_only_content_without_the_legacy_ocr_path():
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    from tests.test_images import vector_pdf

    writer = PdfWriter()
    page = writer.add_page(PdfReader(io.BytesIO(vector_pdf())).pages[0])
    stream = DecodedStreamObject()
    stream.set_data(b"0 0 1 RG 3 w 50 100 m 150 180 l 250 160 l 400 300 l S")
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)

    class EmptyOCR:
        max_pages = 200

        async def read(self, data, content_type):
            return "", 1

    vision = FakeChat({"description": "An unlabeled blue line chart."})
    client, services, lib, worker = setup(vision)
    worker.ocr = EmptyOCR()
    worker.layout = FakeLayout(caption=False, body=False)
    document = upload(client, lib, "unlabeled.pdf", data=out.getvalue())
    await drain(services, worker)
    stored = await services.documents.get(lib, document)
    assert stored["status"] == "ready" and stored["imageCount"] == 1
    assert stored["extraction"]["method"] == "document_intelligence_layout" and len(vision.calls) == 1


async def test_skipped_visuals_are_disclosed_in_search_coverage():
    client, services, lib, worker = setup(FakeChat({"description": "chart"}))
    worker.image_min_bytes = 100000
    from tests.test_images import office

    upload(client, lib, "charts.docx", data=office("docx"))
    await drain(services, worker)
    result = client.post(
        f"/v1/libraries/{lib}/search", headers={"X-Agent-Id": "agent"}, json={"query": "GDP chart"}
    ).json()
    assert any("1 visual candidates were skipped" in note for note in result["coverage"]["notes"])


async def test_failed_cleanup_send_can_be_retried_after_a_document_is_gone():
    client, services, lib, worker = setup(FakeChat({"description": "chart"}))
    # A late attempt may leave artifacts after the deletion's first purge completed.
    await services.blobs.upload(f"{lib}/deleted/text/1/plan.json", b"{}", "application/json")
    await worker.handle({"type": "ingest", "libraryId": lib, "documentId": "deleted", "generation": 1})
    await drain(services, worker)
    assert not services.blobs.blobs
