"""Small invented chart fixtures for the image retrieval exit check. No real economic data."""

from __future__ import annotations

import io

from scripts.sample_docs.images import INK, canvas, font, title


def margin_png() -> bytes:
    image, draw = canvas(seed=42)
    title(draw, "Regional gross margin, Q4 2025", "Invented demonstration data; percent")
    for i, (region, margin) in enumerate([("EMEA", 28), ("APAC", 34), ("North America", 31)]):
        x = 150 + i * 330
        y = 680 - margin * 12
        draw.rectangle((x, y, x + 170, 680), fill=(45, 110, 180))
        draw.text((x, 705), region, font=font(28), fill=INK)
        draw.text((x + 35, y - 45), f"{margin}%", font=font(36, True), fill=INK)
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def margin_office(kind: str) -> bytes:
    out = io.BytesIO()
    explanation = "EMEA margins weakened because freight costs rose. APAC benefited from a richer product mix."
    if kind == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Regional profitability", 1)
        document.add_paragraph(explanation)
        document.add_picture(io.BytesIO(margin_png()))
        document.add_paragraph("Figure 1. Regional gross margin. Figures are invented for this demonstration.")
    else:
        from pptx import Presentation
        from pptx.util import Inches

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[5])
        slide.shapes.title.text = "Regional profitability"
        slide.shapes.add_picture(io.BytesIO(margin_png()), Inches(1), Inches(1.5), width=Inches(8))
        slide.notes_slide.notes_text_frame.text = explanation
    document.save(out)
    return out.getvalue()


def gdp_pdf() -> bytes:
    """A chart composed entirely of PDF drawing commands, with a supporting paragraph."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font_dict = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_dict})}
    )
    stream = DecodedStreamObject()
    lines = [
        "BT /F1 20 Tf 50 735 Td (Economic outlook) Tj ET",
        "BT /F1 16 Tf 50 690 Td (Real GDP growth over time, 2019-2024) Tj ET",
        "BT /F1 11 Tf 50 665 Td (Invented example economy; annual change in percent) Tj ET",
        "0.8 0.8 0.8 RG 1 w 70 300 m 550 300 l S",
        "0 0 0 RG 2 w 70 220 m 70 590 l S 70 380 m 550 380 l S",
        "BT /F1 11 Tf 30 580 Td (6%) Tj ET",
        "BT /F1 11 Tf 30 380 Td (0%) Tj ET",
        "BT /F1 11 Tf 24 230 Td (-4%) Tj ET",
        "0.1 0.3 0.8 RG 4 w 90 450 m 180 240 l 270 555 l 360 485 l 450 455 l 540 470 l S",
    ]
    for x, year in zip([90, 180, 270, 360, 450, 540], range(2019, 2025), strict=True):
        lines.append(f"BT /F1 11 Tf {x - 12} 195 Td ({year}) Tj ET")
    lines.extend(
        [
            "BT /F1 12 Tf 50 140 Td (The contraction in 2020 was followed by a strong rebound in 2021.) Tj ET",
            "BT /F1 12 Tf 50 120 Td (Investment and productivity supported the subsequent expansion.) Tj ET",
            "BT /F1 10 Tf 50 80 Td (Figure 1. Synthetic data for retrieval testing, not an economic forecast.) Tj ET",
        ]
    )
    stream.set_data("\n".join(lines).encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def fixtures() -> list[tuple[str, bytes, str]]:
    from scripts.sample_docs.images import bar_chart

    out = io.BytesIO()
    bar_chart().save(out, "PNG")
    return [
        ("Economic-outlook.pdf", gdp_pdf(), "application/pdf"),
        (
            "Regional-review.docx",
            margin_office("docx"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (
            "Regional-results.pptx",
            margin_office("pptx"),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
        ("Rainfall.png", out.getvalue(), "image/png"),
    ]
