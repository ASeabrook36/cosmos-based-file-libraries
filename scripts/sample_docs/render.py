"""Renderers shared by the sample-document generator: text-based document types and a minimal PDF writer."""

from __future__ import annotations

import datetime as dt
import html
import textwrap
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Doc:
    """A text document: sections are (heading, body) or (heading, body, table). Paragraphs split on blank lines."""

    slug: str
    title: str
    author: str
    date: str  # YYYY-MM-DD
    sections: list[tuple]
    q: str = ""
    a: str = ""
    extra: dict = field(default_factory=dict)

    @property
    def when(self) -> dt.datetime:
        return dt.datetime.fromisoformat(self.date).replace(hour=10, tzinfo=dt.UTC)


def paras(body: str) -> list[str]:
    return [" ".join(p.split()) for p in body.strip().split("\n\n") if p.strip()]


def ascii_only(s: str) -> str:
    table = {
        "–": "-",
        "—": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "°": " deg",
        "é": "e",
        "è": "e",
        "ü": "u",
        "ö": "o",
        "á": "a",
        "í": "i",
        "ó": "o",
        "ç": "c",
        "ñ": "n",
        "×": "x",
        "→": "->",
        "£": "GBP ",
        "€": "EUR ",
        "…": "...",
        "µ": "u",
        "−": "-",
        "½": "1/2",
    }
    return "".join(table.get(c, c) for c in s).encode("ascii", "replace").decode()


# ---------- Markdown / text / HTML ----------


def to_md(d: Doc) -> str:
    out = [f"# {d.title}", "", f"*{d.author}, {d.date}*", ""]
    for sec in d.sections:
        out += [f"## {sec[0]}", ""]
        for p in paras(sec[1]):
            out += [p, ""]
        if len(sec) > 2:
            hdr, *rows = sec[2]
            out += ["| " + " | ".join(hdr) + " |", "|" + "---|" * len(hdr)]
            out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows] + [""]
    return "\n".join(out)


def to_txt(d: Doc) -> str:
    out = [d.title.upper(), f"By {d.author}  |  {d.date}", ""]
    for sec in d.sections:
        out += [sec[0], "-" * len(sec[0]), ""]
        for p in paras(sec[1]):
            out += textwrap.wrap(p, 88) + [""]
        if len(sec) > 2:
            for r in sec[2]:
                out.append("  " + " | ".join(str(c) for c in r))
            out.append("")
    return "\n".join(out)


def to_html(d: Doc) -> str:
    e = html.escape
    out = [
        "<!doctype html>",
        '<html lang="en">',
        "<head>",
        '<meta charset="utf-8">',
        f"<title>{e(d.title)}</title>",
        f'<meta name="author" content="{e(d.author)}">',
        f'<meta name="date" content="{d.date}">',
        "</head>",
        "<body>",
        f"<h1>{e(d.title)}</h1>",
        f'<p class="byline">{e(d.author)}, <time datetime="{d.date}">{d.date}</time></p>',
    ]
    for sec in d.sections:
        out.append(f"<h2>{e(sec[0])}</h2>")
        out += [f"<p>{e(p)}</p>" for p in paras(sec[1])]
        if len(sec) > 2:
            hdr, *rows = sec[2]
            out.append("<table>")
            out.append("<tr>" + "".join(f"<th>{e(str(c))}</th>" for c in hdr) + "</tr>")
            out += ["<tr>" + "".join(f"<td>{e(str(c))}</td>" for c in r) + "</tr>" for r in rows]
            out.append("</table>")
    out += ["</body>", "</html>"]
    return "\n".join(out)


# ---------- DOCX ----------


def to_docx(d: Doc, path: Path, picture: Path | None = None) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.core_properties.author = d.author
    doc.core_properties.title = d.title
    doc.core_properties.created = d.when
    doc.core_properties.modified = d.when
    doc.add_heading(d.title, level=0)
    for i, sec in enumerate(d.sections):
        doc.add_heading(sec[0], level=1)
        for p in paras(sec[1]):
            doc.add_paragraph(p)
        if len(sec) > 2:
            hdr, *rows = sec[2]
            t = doc.add_table(rows=1, cols=len(hdr))
            t.style = "Table Grid"
            for c, h in zip(t.rows[0].cells, hdr, strict=True):
                c.text = str(h)
            for r in rows:
                for c, v in zip(t.add_row().cells, r, strict=True):
                    c.text = str(v)
        if picture and i == 0:
            doc.add_picture(str(picture), width=Inches(5.5))
    doc.save(path)


# ---------- PDF (hand-written: Helvetica, text layer, Info dictionary) ----------


def _esc(s: str) -> str:
    return ascii_only(s).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def to_pdf(d: Doc, path: Path) -> None:
    W, H, M = 612, 792, 64
    pages: list[list[str]] = []
    cur: list[str] = []
    y = H - M

    def newpage():
        nonlocal cur, y
        if cur:
            pages.append(cur)
        cur, y = [], H - M

    def line(text: str, size: int, bold: bool = False, gap: float = 1.35):
        nonlocal y
        if y - size * gap < M:
            newpage()
        font = "F2" if bold else "F1"
        cur.append(f"BT /{font} {size} Tf {M} {y:.1f} Td ({_esc(text)}) Tj ET")
        y -= size * gap

    def wrapped(text: str, size: int, bold: bool = False):
        for ln in textwrap.wrap(ascii_only(text), int((W - 2 * M) / (size * 0.5))):
            line(ln, size, bold)

    wrapped(d.title, 20, True)
    y -= 4
    line(f"{d.author}  -  {d.date}", 10)
    y -= 10
    for sec in d.sections:
        y -= 6
        wrapped(sec[0], 14, True)
        for p in paras(sec[1]):
            wrapped(p, 11)
            y -= 5
        if len(sec) > 2:
            for r in sec[2]:
                line("   ".join(str(c) for c in r), 10)
            y -= 5
    newpage()

    objs: list[bytes] = []

    def add(b: str | bytes) -> int:
        objs.append(b.encode("latin-1") if isinstance(b, str) else b)
        return len(objs)

    add("<< /Type /Catalog /Pages 2 0 R >>")  # 1
    kids = " ".join(f"{5 + 2 * i} 0 R" for i in range(len(pages)))
    add(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>")  # 2
    add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")  # 3
    add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")  # 4
    for i, ops in enumerate(pages):
        content = "\n".join(ops).encode("latin-1")
        add(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {W} {H}] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> "
            f"/Contents {6 + 2 * i} 0 R >>"
        )
        add(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    stamp = d.when.strftime("D:%Y%m%d%H%M%S+00'00'")
    info = add(f"<< /Title ({_esc(d.title)}) /Author ({_esc(d.author)}) /CreationDate ({stamp}) /ModDate ({stamp}) >>")

    buf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, o in enumerate(objs, 1):
        offsets.append(len(buf))
        buf += f"{n} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(buf)
    buf += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    buf += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    buf += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info {info} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(buf))
